"""Parser para planilhas de Plano de Ação (PA XX) da ARPE.

Formato esperado:
- Aba "Mapa Estratégico": ignorada.
- Abas "PA 01", "PA 02", ...: cada uma contém UMA iniciativa com metadados
  no topo, seguida de uma tabela de KRs (Nº inteiro) e etapas (Nº decimal).

Estrutura da aba:
  - Metadados em blocos label:valor (busca por label, não por posição).
  - Cabeçalho da tabela: linha com "Nº" e "ETAPAS DO PLANO DE AÇÃO".
  - Cada linha de dados: Nº em col B.
    - Nº inteiro (1, 2, 3) = KR
    - Nº decimal (1.1, 1.2) = etapa do KR pai
  - 3 blocos de justificativa (J-P, Q-W, X-AD), cada um com 7 colunas.
"""

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.domain.import_spreadsheet.schemas import (
    ImportIssue,
    JustificativaData,
    ValidatedRow,
)

_PA_SHEET_PATTERN = re.compile(r"^PA[\s_-]*\d+$", re.IGNORECASE)

_METADATA_LABELS: dict[str, str] = {
    "PERSPECTIVA ESTRATÉGICA": "perspective",
    "OBJETIVO ESTRATÉGICO": "objective",
    "SETOR RESPONSÁVEL": "responsible_sector",
    "NOME DA INICIATIVA ESTRATÉGICA": "initiative_name",
    "GESTOR RESPONSÁVEL": "responsible_manager",
    "DESCRIÇÃO RESUMIDA DA INICIATIVA ESTRATÉGICA": "summary",
    "VALOR ESTIMADO TOTAL": "total_estimated_value",
    "VALOR ESTIMADO 2026": "current_year_estimated_value",
}

_REQUIRED_METADATA = {
    "perspective": "Perspectiva Estratégica",
    "objective": "Objetivo Estratégico",
    "responsible_sector": "Setor Responsável",
    "initiative_name": "Nome da Iniciativa",
    "responsible_manager": "Gestor Responsável",
    "summary": "Descrição Resumida",
}

_STATUS_MAP = {
    "a iniciar": "Not Started",
    "not started": "Not Started",
    "em andamento": "In Progress",
    "in progress": "In Progress",
    "concluido": "Completed",
    "concluído": "Completed",
    "completed": "Completed",
    "atrasado": "Delayed",
    "delayed": "Delayed",
}

_JUSTIFICATIVA_BASE_COLS = (10, 17, 24)  # colunas J, Q, X (1-indexed)

_DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%y")


def _text(raw: Any) -> str:
    if raw is None:
        return ""
    return str(raw).strip()


def _normalize_label(raw: Any) -> str:
    return _text(raw).upper().rstrip(":").strip()


def _parse_date(raw: Any) -> date | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_status(raw: Any) -> str | None:
    return _STATUS_MAP.get(_text(raw).lower())


def _find_metadata_value(ws, label: str) -> str | None:
    """Procura `label` em qualquer célula e retorna o próximo valor não-vazio à direita."""
    target = _normalize_label(label)
    for row in ws.iter_rows():
        for cell in row:
            if _normalize_label(cell.value) == target:
                for col_offset in range(1, 10):
                    val = ws.cell(row=cell.row, column=cell.column + col_offset).value
                    text = _text(val)
                    if text:
                        return text
                return None
    return None


def _find_header_row(ws) -> int | None:
    """Encontra a linha do cabeçalho buscando 'ETAPAS DO PLANO DE AÇÃO'."""
    for row in ws.iter_rows(min_col=3, max_col=3):
        for cell in row:
            if "ETAPAS DO PLANO DE AÇÃO" in _normalize_label(cell.value):
                return cell.row
    return None


def _parse_justificativa(ws, row_num: int, base_col: int) -> JustificativaData | None:
    """Extrai um bloco de justificativa (7 colunas) a partir de `base_col`."""
    problema = _text(ws.cell(row=row_num, column=base_col).value)
    if not problema:
        return None

    return JustificativaData(
        problema=problema,
        data_registro=_parse_date(ws.cell(row=row_num, column=base_col + 1).value),
        acao_corretiva=_text(ws.cell(row=row_num, column=base_col + 2).value),
        responsavel=_text(ws.cell(row=row_num, column=base_col + 3).value),
        previsao_solucao=_parse_date(ws.cell(row=row_num, column=base_col + 4).value),
        sob_governanca_arpe=_text(ws.cell(row=row_num, column=base_col + 5).value).lower()
        in ("sim", "s", "yes", "true", "1"),
        resolvido=_text(ws.cell(row=row_num, column=base_col + 6).value).lower()
        in ("sim", "s", "yes", "true", "1"),
    )


def _parse_pa_sheet(ws, sheet_name: str) -> tuple[list[ValidatedRow], list[ImportIssue]]:
    issues: list[ImportIssue] = []

    # 1) Metadados
    metadata: dict[str, str] = {}
    for label, key in _METADATA_LABELS.items():
        metadata[key] = _find_metadata_value(ws, label) or ""

    missing = [label for key, label in _REQUIRED_METADATA.items() if not metadata.get(key)]
    if missing:
        for key, label in _REQUIRED_METADATA.items():
            if not metadata.get(key):
                issues.append(ImportIssue(
                    line=1,
                    field=key,
                    message=f"[{sheet_name}] Metadado obrigatório ausente: {label}",
                ))
        return [], issues

    # 2) Cabeçalho
    header_row = _find_header_row(ws)
    if not header_row:
        issues.append(ImportIssue(
            line=1,
            message=f"[{sheet_name}] Cabeçalho 'ETAPAS DO PLANO DE AÇÃO' não encontrado.",
        ))
        return [], issues

    # 3) Iterar linhas de dados
    valid: list[ValidatedRow] = []
    current_kr_code: str | None = None
    current_kr_name: str | None = None

    for row_num in range(header_row + 1, ws.max_row + 1):
        numero_str = _text(ws.cell(row=row_num, column=2).value)
        if not numero_str:
            continue

        try:
            numero_float = float(numero_str.replace(",", "."))
        except ValueError:
            # Linha de sub-cabeçalho ou lixo — ignora
            continue

        name = _text(ws.cell(row=row_num, column=3).value)
        planned_start = _parse_date(ws.cell(row=row_num, column=4).value)
        planned_end = _parse_date(ws.cell(row=row_num, column=5).value)
        actual_start = _parse_date(ws.cell(row=row_num, column=6).value)
        actual_end = _parse_date(ws.cell(row=row_num, column=7).value)
        responsible = _text(ws.cell(row=row_num, column=8).value)
        status_raw = _text(ws.cell(row=row_num, column=9).value)

        # KR (inteiro) vs etapa (decimal)
        if "." not in numero_str:
            kr_num = int(numero_float)
            if not name:
                issues.append(ImportIssue(
                    line=row_num,
                    message=f"[{sheet_name}] KR Nº {kr_num} sem nome.",
                ))
                continue
            current_kr_code = f"KR-{kr_num}"
            current_kr_name = name
            continue

        # Etapa
        if not current_kr_code or not current_kr_name:
            issues.append(ImportIssue(
                line=row_num,
                message=f"[{sheet_name}] Etapa {numero_str} sem KR pai.",
            ))
            continue

        step_num = int(numero_str.split(".")[-1])
        status = _parse_status(status_raw)

        problems: list[str] = []
        if not name:
            problems.append("nome vazio")
        if not planned_start:
            problems.append("data de início previsto ausente")
        if not planned_end:
            problems.append("data de término previsto ausente")
        if not responsible:
            problems.append("responsável vazio")
        if not status:
            problems.append(f"status inválido ({status_raw!r})")

        if problems:
            issues.append(ImportIssue(
                line=row_num,
                field="step",
                message=f"[{sheet_name}] Etapa {numero_str}: " + "; ".join(problems),
            ))
            continue

        justificativas = [
            j for base_col in _JUSTIFICATIVA_BASE_COLS
            if (j := _parse_justificativa(ws, row_num, base_col)) is not None
        ]

        valid.append(ValidatedRow(
            perspective=metadata["perspective"],
            objective=metadata["objective"],
            responsible_sector=metadata["responsible_sector"],
            initiative_name=metadata["initiative_name"],
            responsible_manager=metadata["responsible_manager"],
            summary=metadata["summary"],
            total_estimated_value=metadata.get("total_estimated_value") or None,
            current_year_estimated_value=metadata.get("current_year_estimated_value") or None,
            kr_code=current_kr_code,
            kr_name=current_kr_name,
            kr_goal=current_kr_name,
            sequential_number=step_num,
            step_name=name,
            planned_start_date=planned_start,
            planned_end_date=planned_end,
            actual_start_date=actual_start,
            actual_end_date=actual_end,
            responsible=responsible,
            status=status,
            justificativas=justificativas,
        ))

    return valid, issues


def parse_workbook(content: bytes) -> tuple[list[ValidatedRow], list[ImportIssue], dict[str, Any]]:
    """Parsea o arquivo .xlsx inteiro. Retorna (rows, issues, summary)."""
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        return [], [ImportIssue(line=0, message=f"Falha ao abrir arquivo: {exc}")], {
            "processed_sheets": [],
            "ignored_sheets": [],
        }

    all_sheets = wb.sheetnames
    pa_sheets = [s for s in all_sheets if _PA_SHEET_PATTERN.match(s.strip())]
    ignored_sheets = [s for s in all_sheets if s not in pa_sheets]

    if not pa_sheets:
        return [], [ImportIssue(
            line=0,
            message=(
                "Nenhuma aba 'PA XX' encontrada. "
                f"Abas disponíveis: {', '.join(all_sheets) or '(nenhuma)'}"
            ),
        )], {"processed_sheets": [], "ignored_sheets": ignored_sheets}

    all_valid: list[ValidatedRow] = []
    all_issues: list[ImportIssue] = []

    for sheet_name in pa_sheets:
        valid, issues = _parse_pa_sheet(wb[sheet_name], sheet_name)
        all_valid.extend(valid)
        all_issues.extend(issues)

    return all_valid, all_issues, {
        "processed_sheets": pa_sheets,
        "ignored_sheets": ignored_sheets,
    }